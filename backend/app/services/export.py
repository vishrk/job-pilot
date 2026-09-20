"""PDF/DOCX export. ATS-parseable by construction: plain paragraphs and headings
only — no tables, text boxes, multi-column layout, or images (§7 Phase 2 accept
criteria)."""

import io

import docx
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def _resume_lines(content: dict) -> list[tuple[str, str]]:
    """Returns (style, text) pairs: style is 'h1'/'h2'/'body'."""
    lines = [("h1", "Resume"), ("body", content["summary"])]
    for section in content["sections"]:
        lines.append(("h2", section["experience_id"]))
        for bullet in section["bullets"]:
            lines.append(("body", f"• {bullet['text']}"))
    if content.get("skills"):
        lines.append(("h2", "Skills"))
        lines.append(("body", ", ".join(s["canonical"] for s in content["skills"])))
    return lines


def _cover_letter_lines(content: dict) -> list[tuple[str, str]]:
    lines = [("body", content["greeting"])]
    for p in content["paragraphs"]:
        lines.append(("body", p["text"]))
    lines.append(("body", content["closing"]))
    return lines


def _lines_for(content: dict, kind: str) -> list[tuple[str, str]]:
    return _resume_lines(content) if kind == "resume" else _cover_letter_lines(content)


def to_docx(content: dict, kind: str) -> bytes:
    document = docx.Document()
    for style, text in _lines_for(content, kind):
        heading_level = {"h1": 1, "h2": 2}.get(style)
        if heading_level:
            document.add_heading(text, level=heading_level)
        else:
            document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def to_pdf(content: dict, kind: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER)
    styles = getSampleStyleSheet()
    flowables = []
    for style, text in _lines_for(content, kind):
        style_name = {"h1": "Heading1", "h2": "Heading2", "body": "BodyText"}[style]
        flowables.append(Paragraph(text, styles[style_name]))
        flowables.append(Spacer(1, 6))
    doc.build(flowables)
    return buf.getvalue()
